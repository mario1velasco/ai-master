#!/usr/bin/env python3
"""Create an editable blank PPTX seed using only python-pptx."""
from __future__ import annotations

import argparse
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches


def create_base_theme(output: Path, width_in: float = 13.333333, height_in: float = 7.5) -> Path:
    if width_in <= 1 or height_in <= 1:
        raise ValueError("Slide dimensions must be greater than one inch")
    output.parent.mkdir(parents=True, exist_ok=True)
    presentation = Presentation()
    presentation.slide_width = Inches(width_in)
    presentation.slide_height = Inches(height_in)
    presentation.slides.add_slide(presentation.slide_layouts[6])
    presentation.core_properties.title = "Presentations Open blank seed"
    presentation.core_properties.subject = "Blank slide seed; theme fonts and colors are applied by compile_deck.py"
    presentation.save(str(output))
    return output


def main() -> int:
    parser = argparse.ArgumentParser(description="Create the blank PPTX seed consumed by compile_deck.py")
    parser.add_argument("--output", type=Path, default=Path(__file__).resolve().parent / "base_theme.pptx")
    parser.add_argument("--width-in", type=float, default=13.333333)
    parser.add_argument("--height-in", type=float, default=7.5)
    args = parser.parse_args()
    try:
        path = create_base_theme(args.output.resolve(), args.width_in, args.height_in)
    except (OSError, ValueError) as exc:
        parser.error(str(exc))
        return 2
    print("Created blank seed: {}".format(path))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
