#!/usr/bin/env python3
"""Report PPTX geometry and editable objects using python-pptx public APIs."""
from __future__ import annotations

import argparse
import hashlib
import json
import sys
from pathlib import Path
from typing import Any, Dict, List, Optional

from pptx import Presentation

EMU_PER_INCH = 914400.0


def _color(value: Any) -> Optional[str]:
    try:
        rgb = value.rgb
        return str(rgb) if rgb is not None else None
    except (AttributeError, TypeError, ValueError):
        return None


def _font_data(paragraphs: Any) -> List[Dict[str, Any]]:
    result = []
    for paragraph in paragraphs:
        for run in paragraph.runs:
            result.append({
                "text": run.text,
                "name": run.font.name,
                "size_pt": run.font.size.pt if run.font.size is not None else None,
                "bold": run.font.bold,
                "italic": run.font.italic,
                "color": _color(run.font.color),
                "hyperlink": run.hyperlink.address or None,
            })
    return result


def _shape_report(shape: Any) -> Dict[str, Any]:
    item: Dict[str, Any] = {
        "name": shape.name,
        "type": getattr(shape.shape_type, "name", str(shape.shape_type)),
        "box_in": {
            "x": round(float(shape.left) / EMU_PER_INCH, 4),
            "y": round(float(shape.top) / EMU_PER_INCH, 4),
            "width": round(float(shape.width) / EMU_PER_INCH, 4),
            "height": round(float(shape.height) / EMU_PER_INCH, 4),
        },
    }
    if getattr(shape, "is_placeholder", False):
        item["placeholder_idx"] = shape.placeholder_format.idx
        item["placeholder_type"] = str(shape.placeholder_format.type)
    if getattr(shape, "has_text_frame", False):
        item["text"] = shape.text_frame.text
        item["runs"] = _font_data(shape.text_frame.paragraphs)
    if getattr(shape, "has_table", False):
        table = shape.table
        item["table"] = [[cell.text for cell in row.cells] for row in table.rows]
        item["column_widths_in"] = [round(float(column.width) / EMU_PER_INCH, 4) for column in table.columns]
        item["row_heights_in"] = [round(float(row.height) / EMU_PER_INCH, 4) for row in table.rows]
    if getattr(shape, "has_chart", False):
        chart = shape.chart
        item["chart"] = {
            "type": str(chart.chart_type),
            "title": chart.chart_title.text_frame.text if chart.has_title else None,
            "has_legend": chart.has_legend,
            "series": [{"name": series.name, "values": list(series.values)} for series in chart.series],
        }
        try:
            item["chart"]["categories"] = [str(category.label) for category in chart.plots[0].categories]
        except (AttributeError, TypeError, ValueError):
            item["chart"]["categories"] = []
    if getattr(shape.shape_type, "name", "") in ("PICTURE", "LINKED_PICTURE"):
        image = shape.image
        blob = image.blob
        item["image"] = {
            "content_type": image.content_type,
            "width_px": image.size[0],
            "height_px": image.size[1],
            "sha256": hashlib.sha256(blob).hexdigest(),
        }
    try:
        item["fill_rgb"] = _color(shape.fill.fore_color)
    except (AttributeError, TypeError, ValueError):
        item["fill_rgb"] = None
    try:
        item["line_rgb"] = _color(shape.line.color)
    except (AttributeError, TypeError, ValueError):
        item["line_rgb"] = None
    return item


def inspect(path: Path) -> Dict[str, Any]:
    presentation = Presentation(str(path))
    slides = []
    for number, slide in enumerate(presentation.slides, start=1):
        notes = None
        if slide.has_notes_slide:
            frame = slide.notes_slide.notes_text_frame
            notes = frame.text if frame is not None else None
        slides.append({
            "number": number,
            "layout": slide.slide_layout.name,
            "notes": notes,
            "shapes": [_shape_report(shape) for shape in slide.shapes],
        })
    return {
        "file": str(path.resolve()),
        "slide_count": len(slides),
        "slide_size_in": {
            "width": round(float(presentation.slide_width) / EMU_PER_INCH, 4),
            "height": round(float(presentation.slide_height) / EMU_PER_INCH, 4),
        },
        "metadata": {
            "title": presentation.core_properties.title,
            "subject": presentation.core_properties.subject,
            "author": presentation.core_properties.author,
            "keywords": presentation.core_properties.keywords,
        },
        "masters": [{"name": master.name, "layout_names": [layout.name for layout in master.slide_layouts]} for master in presentation.slide_masters],
        "slides": slides,
    }


def main() -> int:
    parser = argparse.ArgumentParser(description="Inspect slides, notes, typography, tables, charts, images, and template layouts")
    parser.add_argument("presentation", type=Path)
    parser.add_argument("--output", type=Path, help="Write JSON report to this path; stdout when omitted")
    args = parser.parse_args()
    path = args.presentation.resolve()
    if not path.is_file():
        parser.error("PPTX not found: {}".format(path))
    try:
        report = json.dumps(inspect(path), ensure_ascii=False, indent=2)
    except Exception as exc:
        print("ERROR: {}".format(exc), file=sys.stderr)
        return 1
    if args.output:
        output = args.output.resolve()
        output.parent.mkdir(parents=True, exist_ok=True)
        output.write_text(report + "\n", encoding="utf-8")
        print("Wrote inspection: {}".format(output))
    else:
        print(report)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
