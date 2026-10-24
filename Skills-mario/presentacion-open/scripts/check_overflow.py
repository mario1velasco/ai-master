#!/usr/bin/env python3
"""Validate PPTX geometry, text containers, and required native objects."""
from __future__ import annotations

import argparse
import json
import math
import sys
from pathlib import Path
from typing import Any, Dict, Iterable, List, Mapping, Optional, Sequence, Set, Tuple

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

EMU_PER_INCH = 914400.0
DEFAULT_MARGIN_IN = 0.6
MAX_CONTAINER_CHARS = 180
OVERLAP_TOLERANCE_IN = 0.001


def label(shape: Any) -> str:
    return str(getattr(shape, "name", "") or "unnamed shape")


def _is_line(shape: Any) -> bool:
    try:
        return shape.shape_type == MSO_SHAPE_TYPE.LINE
    except Exception:
        return False


def _shape_list(slide: Any) -> List[Any]:
    # GroupShape children use group-relative coordinates; inspect the group's
    # own bounding box and leave child geometry to the visual review.
    return list(slide.shapes)


def _box(shape: Any) -> Tuple[float, float, float, float]:
    x = float(shape.left) / EMU_PER_INCH
    y = float(shape.top) / EMU_PER_INCH
    width = float(shape.width) / EMU_PER_INCH
    height = float(shape.height) / EMU_PER_INCH
    return x, y, x + width, y + height


def _shape_text(shape: Any) -> str:
    if getattr(shape, "has_text_frame", False):
        return str(shape.text_frame.text or "")
    return ""


def _geometry_issues(slide: Any, slide_number: int, slide_width: float, slide_height: float, margin: float) -> List[str]:
    issues: List[str] = []
    right_limit = slide_width - margin
    bottom_limit = slide_height - margin
    for shape in _shape_list(slide):
        x, y, right, bottom = _box(shape)
        faults = []
        if x < margin - 1e-6:
            faults.append("left {:.3f} < {:.3f}".format(x, margin))
        if y < margin - 1e-6:
            faults.append("top {:.3f} < {:.3f}".format(y, margin))
        if right > right_limit + 1e-6:
            faults.append("right {:.3f} > {:.3f}".format(right, right_limit))
        if bottom > bottom_limit + 1e-6:
            faults.append("bottom {:.3f} > {:.3f}".format(bottom, bottom_limit))
        if faults:
            issues.append("slide {}: {} outside safe area ({})".format(slide_number, label(shape), "; ".join(faults)))
    return issues


def _text_issues(slide: Any, slide_number: int, max_chars: int) -> List[str]:
    issues: List[str] = []
    for shape in _shape_list(slide):
        text = _shape_text(shape)
        if text and len(text.strip()) > max_chars:
            issues.append("slide {}: {} has {} characters (maximum {})".format(slide_number, label(shape), len(text.strip()), max_chars))
        if getattr(shape, "has_table", False):
            for row_index, row in enumerate(shape.table.rows):
                for col_index, cell in enumerate(row.cells):
                    cell_text = (cell.text or "").strip()
                    if len(cell_text) > max_chars:
                        issues.append("slide {}: table {} cell ({},{}) has {} characters (maximum {})".format(slide_number, label(shape), row_index + 1, col_index + 1, len(cell_text), max_chars))
    return issues


def _allowed_pair(a: str, b: str, allowed: Set[Tuple[str, str]]) -> bool:
    key = tuple(sorted((a.casefold(), b.casefold())))
    return key in allowed


def _overlap_issues(slide: Any, slide_number: int, allowed: Set[Tuple[str, str]]) -> List[str]:
    issues: List[str] = []
    shapes = _shape_list(slide)
    for i, first in enumerate(shapes):
        if _is_line(first):
            continue
        ax1, ay1, ax2, ay2 = _box(first)
        for second in shapes[i + 1 :]:
            if _is_line(second) or _allowed_pair(label(first), label(second), allowed):
                continue
            bx1, by1, bx2, by2 = _box(second)
            overlap_w = min(ax2, bx2) - max(ax1, bx1)
            overlap_h = min(ay2, by2) - max(ay1, by1)
            if overlap_w > OVERLAP_TOLERANCE_IN and overlap_h > OVERLAP_TOLERANCE_IN:
                issues.append("slide {}: {} overlaps {} by {:.3f} x {:.3f} in".format(slide_number, label(first), label(second), overlap_w, overlap_h))
    return issues


def _object_requirements(presentation: Any, requirements: Mapping[str, Any]) -> List[str]:
    issues: List[str] = []
    slide_count = len(presentation.slides)
    expected = requirements.get("total_slides")
    if expected is not None and slide_count != int(expected):
        issues.append("deck has {} slides; requirements.total_slides is {}".format(slide_count, expected))
    for key, predicate, noun in (
        ("required_native_table_slides", "has_table", "native table"),
        ("required_native_chart_slides", "has_chart", "native chart"),
    ):
        requested = requirements.get(key, [])
        if not isinstance(requested, list):
            issues.append("requirements.{} must be a list".format(key))
            continue
        for number in requested:
            index = int(number) - 1
            if index < 0 or index >= slide_count:
                issues.append("requirements.{} references missing slide {}".format(key, number))
                continue
            found = any(bool(getattr(shape, predicate, False)) for shape in presentation.slides[index].shapes)
            if not found:
                issues.append("slide {} is required to contain a {}".format(number, noun))
    return issues


def _chart_issues(presentation: Any) -> List[str]:
    issues: List[str] = []
    for slide_number, slide in enumerate(presentation.slides, start=1):
        for shape in slide.shapes:
            if not getattr(shape, "has_chart", False):
                continue
            chart = shape.chart
            if chart.has_title and not (chart.chart_title.text_frame.text or "").strip():
                issues.append("slide {}: {} has an empty chart title placeholder".format(slide_number, label(shape)))
            for series in chart.series:
                for field in ("values", "x_values", "bubble_sizes"):
                    values = getattr(series, field, None)
                    if values is None:
                        continue
                    for value in values:
                        try:
                            number = float(value)
                            if not math.isfinite(number):
                                raise ValueError()
                        except (TypeError, ValueError):
                            issues.append("slide {}: {} contains a non-finite chart {} value".format(slide_number, label(shape), field))
                            break
    return issues


def _font_policy_issues(presentation: Any, spec: Mapping[str, Any]) -> List[str]:
    policy = spec.get("font_policy")
    if not isinstance(policy, dict) or not policy.get("families"):
        return []
    allowed = {str(value).casefold() for value in policy["families"]}
    issues: List[str] = []
    for slide_number, slide in enumerate(presentation.slides, start=1):
        for shape in _shape_list(slide):
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    for cell in row.cells:
                        for paragraph in cell.text_frame.paragraphs:
                            for run in paragraph.runs:
                                name = run.font.name
                                if name and name.casefold() not in allowed:
                                    issues.append("slide {}: {} uses font '{}' outside font_policy.families".format(slide_number, label(shape), name))
            if getattr(shape, "has_chart", False):
                name = shape.chart.font.name
                if name and name.casefold() not in allowed:
                    issues.append("slide {}: chart {} uses font '{}' outside font_policy.families".format(slide_number, label(shape), name))
            if not getattr(shape, "has_text_frame", False):
                continue
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    name = run.font.name
                    if name and name.casefold() not in allowed:
                        issues.append("slide {}: {} uses font '{}' outside font_policy.families".format(slide_number, label(shape), name))
    return issues


def inspect_presentation(
    presentation_or_path: Any,
    margin_in: Optional[float] = None,
    requirements: Optional[Mapping[str, Any]] = None,
    spec: Optional[Mapping[str, Any]] = None,
    allowed_overlaps: Optional[Iterable[Sequence[str]]] = None,
    selected_slides: Optional[Set[int]] = None,
    check_overlaps: bool = True,
) -> List[str]:
    presentation = Presentation(str(presentation_or_path)) if isinstance(presentation_or_path, (str, Path)) else presentation_or_path
    if margin_in is None:
        theme = (spec or {}).get("theme", {})
        margin_in = float(theme.get("safe_margin_inches", DEFAULT_MARGIN_IN)) if isinstance(theme, dict) else DEFAULT_MARGIN_IN
    if requirements is None:
        requirements = (spec or {}).get("requirements", {})
    if not isinstance(requirements, Mapping):
        requirements = {}
    allowed: Set[Tuple[str, str]] = set()
    for pair in allowed_overlaps or requirements.get("allowed_overlaps", []):
        if isinstance(pair, (list, tuple)) and len(pair) == 2:
            allowed.add(tuple(sorted((str(pair[0]).casefold(), str(pair[1]).casefold()))))
    width = float(presentation.slide_width) / EMU_PER_INCH
    height = float(presentation.slide_height) / EMU_PER_INCH
    issues: List[str] = []
    for slide_number, slide in enumerate(presentation.slides, start=1):
        issues.extend(_text_issues(slide, slide_number, MAX_CONTAINER_CHARS))
        if selected_slides is None or slide_number in selected_slides:
            issues.extend(_geometry_issues(slide, slide_number, width, height, float(margin_in)))
            if check_overlaps:
                issues.extend(_overlap_issues(slide, slide_number, allowed))
    issues.extend(_object_requirements(presentation, requirements))
    issues.extend(_chart_issues(presentation))
    issues.extend(_font_policy_issues(presentation, spec or {}))
    return issues


def _load_spec(path: Optional[Path]) -> Mapping[str, Any]:
    if path is None:
        return {}
    with path.open("r", encoding="utf-8") as handle:
        data = json.load(handle)
    if not isinstance(data, dict):
        raise ValueError("deck spec root must be an object")
    return data


def main() -> int:
    parser = argparse.ArgumentParser(description="Validate safe margins, overlaps, text lengths, charts, tables, and deck requirements")
    parser.add_argument("presentation", type=Path, help="PPTX to inspect")
    parser.add_argument("--spec", type=Path, help="Optional deck.json for requirements and theme margin")
    parser.add_argument("--margin", type=float, default=None, help="Safe margin in inches; overrides deck.json")
    parser.add_argument("--skip-slide", type=int, action="append", default=[], help="Skip geometry and overlap checks for a preserved template slide (one-based); text and requirements are still checked")
    parser.add_argument("--allow-overlap", nargs=2, action="append", default=[], metavar=("SHAPE_A", "SHAPE_B"), help="Permit a named intentional overlap pair")
    parser.add_argument("--no-overlap-check", action="store_true", help="Skip shape overlap detection")
    args = parser.parse_args()
    pptx_path = args.presentation.resolve()
    if not pptx_path.is_file():
        print("ERROR: PPTX not found: {}".format(pptx_path), file=sys.stderr)
        return 1
    try:
        spec = _load_spec(args.spec.resolve() if args.spec else None)
        presentation = Presentation(str(pptx_path))
        skip = set(args.skip_slide)
        selected = {i for i in range(1, len(presentation.slides) + 1) if i not in skip}
        issues = inspect_presentation(presentation, args.margin, spec=spec, allowed_overlaps=args.allow_overlap, selected_slides=selected if args.skip_slide else None, check_overlaps=not args.no_overlap_check)
    except Exception as error:
        print("ERROR: could not inspect {}: {}".format(pptx_path, error), file=sys.stderr)
        return 1
    if issues:
        print("PPTX validation issues in {}:".format(pptx_path))
        for issue in issues:
            print("- " + issue)
        return 1
    print("OK: {} passed the requested geometry and object checks.".format(pptx_path))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
