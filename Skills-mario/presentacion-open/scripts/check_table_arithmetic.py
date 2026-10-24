#!/usr/bin/env python3
"""Check only user-declared table sum contracts in a PPTX."""
from __future__ import annotations

import argparse
import json
import re
import sys
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Dict, List, Mapping, Sequence, Tuple

from pptx import Presentation


def parse_displayed_number(value: str, decimal_separator: str, thousands_separator: str, ignored: Sequence[str]) -> Decimal:
    text = str(value).strip()
    negative = text.startswith("(") and text.endswith(")")
    if negative:
        text = text[1:-1]
    for token in ignored:
        text = text.replace(str(token), "")
    if thousands_separator:
        text = text.replace(thousands_separator, "")
    if decimal_separator != ".":
        text = text.replace(decimal_separator, ".")
    normalized = re.sub(r"[^0-9.+-]", "", text)
    if normalized.count(".") > 1 or normalized.count("+") > 1 or normalized.count("-") > 1:
        raise ValueError("ambiguous numeric value '{}'".format(value))
    try:
        number = Decimal(normalized)
    except InvalidOperation:
        raise ValueError("cannot parse numeric value '{}'".format(value))
    return -number if negative else number


def _cell(table: Any, coordinate: Any, location: str) -> str:
    if not isinstance(coordinate, list) or len(coordinate) != 2:
        raise ValueError("{} must be [row, column] using zero-based indices".format(location))
    row, column = int(coordinate[0]), int(coordinate[1])
    if row < 0 or column < 0 or row >= len(table.rows) or column >= len(table.columns):
        raise ValueError("{} is outside the table".format(location))
    return table.cell(row, column).text


def check(pptx_path: Path, spec_path: Path) -> List[str]:
    spec = json.loads(spec_path.read_text(encoding="utf-8"))
    contracts = spec.get("requirements", {}).get("arithmetic_contracts", [])
    if not isinstance(contracts, list):
        return ["requirements.arithmetic_contracts must be a list"]
    if not contracts:
        return []
    presentation = Presentation(str(pptx_path))
    issues: List[str] = []
    for index, contract in enumerate(contracts, start=1):
        location = "arithmetic contract {}".format(index)
        if not isinstance(contract, dict) or contract.get("operation") != "sum":
            issues.append("{} must declare operation='sum'".format(location))
            continue
        slide_number = int(contract.get("slide", 0))
        table_number = int(contract.get("table", 0))
        if slide_number < 1 or slide_number > len(presentation.slides):
            issues.append("{} references missing slide {}".format(location, slide_number))
            continue
        tables = [shape.table for shape in presentation.slides[slide_number - 1].shapes if getattr(shape, "has_table", False)]
        if table_number < 1 or table_number > len(tables):
            issues.append("{} references missing table {} on slide {}".format(location, table_number, slide_number))
            continue
        table = tables[table_number - 1]
        inputs = contract.get("inputs")
        if not isinstance(inputs, list) or not inputs:
            issues.append("{} needs a non-empty inputs list".format(location))
            continue
        decimal_separator = str(contract.get("decimal_separator", "."))
        thousands_separator = str(contract.get("thousands_separator", ","))
        ignored = contract.get("ignored_tokens", ["€", "$", "£", "¥", "%", " ", "\u00a0"])
        if not isinstance(ignored, list):
            issues.append("{}.ignored_tokens must be a list".format(location))
            continue
        try:
            values = [parse_displayed_number(_cell(table, point, location + ".inputs"), decimal_separator, thousands_separator, ignored) for point in inputs]
            actual = parse_displayed_number(_cell(table, contract.get("result"), location + ".result"), decimal_separator, thousands_separator, ignored)
            expected = sum(values, Decimal("0"))
            tolerance = Decimal(str(contract.get("tolerance", "0")))
            if tolerance < 0:
                raise ValueError("tolerance cannot be negative")
            difference = abs(expected - actual)
            if difference > tolerance:
                issues.append("{}: inputs sum to {}, displayed result is {}, difference {} exceeds tolerance {}".format(location, expected, actual, difference, tolerance))
        except (ValueError, InvalidOperation, TypeError, IndexError) as exc:
            issues.append("{}: {}".format(location, exc))
    return issues


def main() -> int:
    parser = argparse.ArgumentParser(description="Check explicit table arithmetic contracts against editable PPTX table cells")
    parser.add_argument("presentation", type=Path)
    parser.add_argument("--spec", type=Path, required=True, help="deck.json containing requirements.arithmetic_contracts")
    args = parser.parse_args()
    pptx_path, spec_path = args.presentation.resolve(), args.spec.resolve()
    if not pptx_path.is_file() or not spec_path.is_file():
        parser.error("PPTX and deck JSON must exist")
    try:
        issues = check(pptx_path, spec_path)
    except Exception as exc:
        print("ERROR: {}".format(exc), file=sys.stderr)
        return 1
    if issues:
        print("Table arithmetic issues:")
        for issue in issues:
            print("- " + issue)
        return 1
    print("OK: all explicit arithmetic contracts passed.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
