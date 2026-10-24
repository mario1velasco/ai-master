#!/usr/bin/env python3
"""Compare preserved PPTX slides and layouts using public python-pptx objects."""
from __future__ import annotations

import argparse
import hashlib
import json
import sys
from pathlib import Path
from typing import Any, Dict, List, Sequence, Set, Tuple

from pptx import Presentation

EMU_PER_INCH = 914400.0


def _safe_color(color_format: Any) -> str:
    try:
        value = color_format.rgb
        return str(value) if value is not None else ""
    except (AttributeError, TypeError, ValueError):
        return ""


def _shape_signature(shape: Any) -> Dict[str, Any]:
    signature: Dict[str, Any] = {
        "name": shape.name,
        "type": getattr(shape.shape_type, "name", str(shape.shape_type)),
        "box": [int(shape.left), int(shape.top), int(shape.width), int(shape.height)],
        "text": shape.text_frame.text if getattr(shape, "has_text_frame", False) else "",
    }
    if getattr(shape, "has_text_frame", False):
        signature["runs"] = [
            (run.text, run.font.name, run.font.size.pt if run.font.size is not None else None, run.font.bold, run.font.italic, _safe_color(run.font.color))
            for paragraph in shape.text_frame.paragraphs for run in paragraph.runs
        ]
    if getattr(shape, "has_table", False):
        signature["table"] = [[cell.text for cell in row.cells] for row in shape.table.rows]
    if getattr(shape, "has_chart", False):
        chart = shape.chart
        signature["chart"] = {
            "type": str(chart.chart_type),
            "title": chart.chart_title.text_frame.text if chart.has_title else None,
            "series": [(series.name, list(series.values)) for series in chart.series],
        }
    if getattr(shape.shape_type, "name", "") in ("PICTURE", "LINKED_PICTURE"):
        signature["image_sha256"] = hashlib.sha256(shape.image.blob).hexdigest()
    try:
        signature["fill"] = _safe_color(shape.fill.fore_color)
    except (AttributeError, TypeError, ValueError):
        signature["fill"] = ""
    try:
        signature["line"] = _safe_color(shape.line.color)
    except (AttributeError, TypeError, ValueError):
        signature["line"] = ""
    return signature


def _slide_signature(slide: Any) -> List[Dict[str, Any]]:
    return [_shape_signature(shape) for shape in slide.shapes]


def compare(source_path: Path, output_path: Path, allowed_changed: Set[int]) -> List[str]:
    source = Presentation(str(source_path))
    output = Presentation(str(output_path))
    issues: List[str] = []
    source_size = (int(source.slide_width), int(source.slide_height))
    output_size = (int(output.slide_width), int(output.slide_height))
    if source_size != output_size:
        issues.append("slide dimensions changed from {} to {} EMU".format(source_size, output_size))
    if len(output.slides) < len(source.slides):
        issues.append("output has {} slides; source template has {}".format(len(output.slides), len(source.slides)))
    if len(source.slide_masters) != len(output.slide_masters):
        issues.append("slide master count changed")
    else:
        for index, (source_master, output_master) in enumerate(zip(source.slide_masters, output.slide_masters), start=1):
            source_names = [layout.name for layout in source_master.slide_layouts]
            output_names = [layout.name for layout in output_master.slide_layouts]
            if source_names != output_names:
                issues.append("slide master {} layout names changed".format(index))
    for number in range(1, min(len(source.slides), len(output.slides)) + 1):
        if number in allowed_changed:
            continue
        if _slide_signature(source.slides[number - 1]) != _slide_signature(output.slides[number - 1]):
            issues.append("preserved template slide {} changed".format(number))
    for number in allowed_changed:
        if number < 1 or number > len(source.slides):
            issues.append("--allow-changed-slide {} does not exist in source".format(number))
    return issues


def main() -> int:
    parser = argparse.ArgumentParser(description="Check slide dimensions, masters/layouts, and unedited source slide objects")
    parser.add_argument("source_template", type=Path)
    parser.add_argument("output_pptx", type=Path)
    parser.add_argument("--allow-changed-slide", type=int, action="append", default=[], help="One-based source slide intentionally edited")
    args = parser.parse_args()
    source, output = args.source_template.resolve(), args.output_pptx.resolve()
    if not source.is_file() or not output.is_file():
        parser.error("Source template and output PPTX must exist")
    try:
        issues = compare(source, output, set(args.allow_changed_slide))
    except Exception as exc:
        print("ERROR: {}".format(exc), file=sys.stderr)
        return 1
    if issues:
        print("Template fidelity issues:")
        for issue in issues:
            print("- " + issue)
        return 1
    print("OK: template dimensions/layouts and all non-exempt source slides match structurally.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
