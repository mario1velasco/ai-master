#!/usr/bin/env python3
"""Convierte un .md a DOCX (Word) sin perder ni añadir contenido. Usa markdown-it-py + python-docx (sin dependencias del sistema).

Mismo Markdown soportado que md_to_pdf.py: títulos #–######, párrafos, negrita/cursiva/código, enlaces (clicables),
listas con viñetas y numeradas (anidadas hasta 3 niveles; cada lista numerada reinicia su numeración), citas,
bloques de código, tablas con cabecera y separadores. Las imágenes se escriben como su texto alternativo y los
comentarios HTML (<!-- … -->) no se escriben. El primer título de nivel 1 es el título del documento.
"""
from __future__ import annotations

import argparse
import sys
from pathlib import Path
from typing import List, Optional

from _md import inline_text, parser

try:
    import docx
    from docx.enum.text import WD_BREAK
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.shared import Pt, RGBColor
except ImportError:
    sys.exit("error: falta python-docx. Instálalo en el venv compartido: python3 scripts/setup_env.py (desde la raíz del repo) y usa .venv/bin/python")

MONO = "Consolas"
LINK_COLOR = RGBColor(0x1D, 0x4E, 0xD8)


class Writer:
    def __init__(self, document):
        self.doc = document
        self.list_stack: List[dict] = []  # {"ordered": bool, "num_id": int|None, "first": bool}

    # ------------------------------------------------------------ numeración
    def new_num(self, style_name: str, start: int) -> int:
        """Crea una numeración nueva basada en la del estilo para que cada lista empiece en su número."""
        numbering = self.doc.part.numbering_part.element
        st = self.doc.styles[style_name].element
        base = st.pPr.numPr.numId.val
        abstract = next(n.abstractNumId.val for n in numbering.num_lst if n.numId == base)
        num = numbering.add_num(abstract)
        override = OxmlElement("w:lvlOverride")
        override.set(qn("w:ilvl"), "0")
        so = OxmlElement("w:startOverride")
        so.set(qn("w:val"), str(start))
        override.append(so)
        num.append(override)
        return num.numId

    @staticmethod
    def set_num(paragraph, num_id: int) -> None:
        pPr = paragraph._p.get_or_add_pPr()
        numPr = OxmlElement("w:numPr")
        ilvl = OxmlElement("w:ilvl")
        ilvl.set(qn("w:val"), "0")
        nid = OxmlElement("w:numId")
        nid.set(qn("w:val"), str(num_id))
        numPr.extend([ilvl, nid])
        pPr.append(numPr)

    # ------------------------------------------------------------ texto en línea
    def add_link(self, paragraph, url: str, text: str, bold: bool, italic: bool):
        rid = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
        run = paragraph.add_run(text)
        run.bold, run.italic, run.underline = bold or None, italic or None, True
        run.font.color.rgb = LINK_COLOR
        hl = OxmlElement("w:hyperlink")
        hl.set(qn("r:id"), rid)
        run._r.addprevious(hl)
        hl.append(run._r)

    def inline(self, paragraph, token, mono: bool = False) -> None:
        bold = italic = False
        href: Optional[str] = None
        for c in token.children or []:
            t = c.type
            if t in ("text", "image", "code_inline"):
                text = c.content
                if href:
                    self.add_link(paragraph, href, text, bold, italic)
                    continue
                run = paragraph.add_run(text)
                run.bold, run.italic = bold or None, italic or None
                if t == "code_inline" or mono:
                    run.font.name = MONO
            elif t == "softbreak":
                paragraph.add_run(" ")
            elif t == "hardbreak":
                paragraph.add_run().add_break(WD_BREAK.LINE)
            elif t == "html_inline":
                if c.content.lower().startswith("<br"):
                    paragraph.add_run().add_break(WD_BREAK.LINE)
                # el resto de HTML en línea (comentarios, <u>…) no aporta texto
            elif t == "strong_open":
                bold = True
            elif t == "strong_close":
                bold = False
            elif t == "em_open":
                italic = True
            elif t == "em_close":
                italic = False
            elif t == "link_open":
                href = c.attrGet("href") or None
            elif t == "link_close":
                href = None

    # ------------------------------------------------------------ bloques
    def paragraph(self, container, style: Optional[str] = None):
        return container.add_paragraph(style=style)

    def list_style(self) -> str:
        level = min(len(self.list_stack), 3)
        top = self.list_stack[-1]
        base = "List Number" if top["ordered"] else "List Bullet"
        if not top["first"]:
            base = "List Continue"
        return base if level == 1 else "{} {}".format(base, level)

    def blocks(self, tokens, i: int, until: Optional[str], container, quote: bool = False) -> int:
        while i < len(tokens):
            t = tokens[i]
            if until and t.type == until:
                return i + 1
            if t.type == "heading_open":
                level = int(t.tag[1])
                p = self.paragraph(container, "Heading {}".format(level))
                self.inline(p, tokens[i + 1])
                i += 3
            elif t.type == "paragraph_open":
                style = None
                if self.list_stack:
                    style = self.list_style()
                elif quote:
                    style = "Quote"
                p = self.paragraph(container, style)
                if self.list_stack:
                    top = self.list_stack[-1]
                    if top["first"] and top["ordered"]:
                        self.set_num(p, top["num_id"])
                    top["first"] = False
                self.inline(p, tokens[i + 1])
                i += 3
            elif t.type in ("bullet_list_open", "ordered_list_open"):
                ordered = t.type == "ordered_list_open"
                entry = {"ordered": ordered, "num_id": None, "first": True}
                self.list_stack.append(entry)
                if ordered:
                    style = "List Number" if len(self.list_stack) == 1 else "List Number {}".format(min(len(self.list_stack), 3))
                    entry["num_id"] = self.new_num(style, int(t.attrGet("start") or 1))
                i = self.blocks(tokens, i + 1, t.type.replace("open", "close"), container, quote)
                self.list_stack.pop()
            elif t.type == "list_item_open":
                self.list_stack[-1]["first"] = True
                i = self.blocks(tokens, i + 1, "list_item_close", container, quote)
            elif t.type == "blockquote_open":
                i = self.blocks(tokens, i + 1, "blockquote_close", container, True)
            elif t.type in ("fence", "code_block"):
                p = self.paragraph(container, "No Spacing")
                lines = t.content.rstrip("\n").split("\n")
                for n, line in enumerate(lines):
                    run = p.add_run(line)
                    run.font.name, run.font.size = MONO, Pt(9)
                    if n < len(lines) - 1:
                        run.add_break(WD_BREAK.LINE)
                i += 1
            elif t.type == "hr":
                p = self.paragraph(container)
                pPr = p._p.get_or_add_pPr()
                bdr = OxmlElement("w:pBdr")
                bottom = OxmlElement("w:bottom")
                for k, v in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", "CBD2D9")):
                    bottom.set(qn(k), v)
                bdr.append(bottom)
                pPr.append(bdr)
                i += 1
            elif t.type == "table_open":
                i = self.table(tokens, i + 1, container)
            else:
                i += 1  # html_block (comentarios) y cierres sueltos
        return i

    def table(self, tokens, i: int, container) -> int:
        rows, row, head_rows = [], [], 0
        while tokens[i].type != "table_close":
            t = tokens[i]
            if t.type in ("th_open", "td_open"):
                row.append((tokens[i + 1], t.type == "th_open"))
                i += 2
            elif t.type == "tr_close":
                rows.append(row)
                row = []
            elif t.type == "thead_close":
                head_rows = len(rows)
            i += 1
        cols = max(len(r) for r in rows)
        tbl = container.add_table(rows=len(rows), cols=cols)
        tbl.style = self.doc.styles["Table Grid"]
        for r, cells in enumerate(rows):
            for c, (inline_tok, is_head) in enumerate(cells):
                p = tbl.cell(r, c).paragraphs[0]
                self.inline(p, inline_tok)
                if is_head:
                    for run in p.runs:
                        run.bold = True
        if head_rows:
            for r in range(head_rows):  # repetir cabecera en cada página
                trPr = tbl.rows[r]._tr.get_or_add_trPr()
                h = OxmlElement("w:tblHeader")
                trPr.append(h)
        container.add_paragraph()  # separación tras la tabla (párrafo vacío, sin texto)
        return i + 1


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("markdown", type=Path)
    ap.add_argument("--output", type=Path, help="por defecto, el mismo nombre con extensión .docx")
    args = ap.parse_args()
    if not args.markdown.is_file():
        sys.exit("error: no existe {}".format(args.markdown))
    out = args.output or args.markdown.with_suffix(".docx")
    md_text = args.markdown.read_text(encoding="utf-8")
    tokens = parser().parse(md_text)

    document = docx.Document()
    normal = document.styles["Normal"]
    normal.font.name, normal.font.size = "Calibri", Pt(11)
    title = next((inline_text(tokens[i + 1]).strip() for i, t in enumerate(tokens) if t.type == "heading_open" and t.tag == "h1"),
                 args.markdown.stem)
    document.core_properties.title = title
    document.core_properties.author = ""
    Writer(document).blocks(tokens, 0, None, document)
    if not any(p.text.strip() for p in document.paragraphs) and not document.tables:
        sys.exit("error: el Markdown no produjo contenido")
    document.save(str(out))
    print("DOCX creado: {} ({} KB)".format(out, out.stat().st_size // 1024))


if __name__ == "__main__":
    main()
