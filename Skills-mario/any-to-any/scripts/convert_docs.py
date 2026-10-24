#!/usr/bin/env python3
"""Extracción automática de cada unidad del inventario a Markdown: trabajo/unidades/<id>_<nombre>/bruto.md (no se edita).

- doc    → markitdown (+ limpieza de rótulos que añade markitdown y no están en el original: «### Notes:», imágenes).
           JSON/XML se envuelven en un bloque de código. PDF con texto: vistas de página en vista/ para comparar el formato.
- sheet  → conversor propio (openpyxl / xlrd / csv): valores tal como se guardan, sin «NaN» ni «10.0» inventados.
- office → LibreOffice (soffice) a docx/pptx/xlsx y después como doc/sheet.
- text   → .md se copia; el resto se escapa para que Markdown lo muestre tal cual (mismas palabras).
- PDF escaneado → páginas PNG (200 ppp) en paginas/ para transcribir con img-to-text; páginas sin texto de un PDF mixto, igual.
- Imágenes incrustadas (docx/pptx/xlsx/epub) → incrustadas/ (markitdown no lee su texto: preguntar si se transcriben).
- Segundo lector: una extracción independiente (pypdf, python-docx, python-pptx, BeautifulSoup) para avisar de texto que
  markitdown haya perdido (OMITIDO?) o añadido (AÑADIDO?). Se guarda en segundo_lector.txt.

Escribe trabajo/unidades/<…>/estado.json y trabajo/conversion.json. Es reanudable (--force para rehacer).
"""
from __future__ import annotations

import argparse
import csv
import datetime as dt
import io
import json
import re
import shutil
import subprocess
import sys
import zipfile
from collections import Counter
from pathlib import Path
from typing import List, Optional

from _md import content_words, escape_line, escape_text, text_words

IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".gif", ".bmp", ".tif", ".tiff", ".webp", ".emf", ".wmf", ".svg"}
SCAN_DPI = 200
PREVIEW_DPI = 80


def slug(s: str) -> str:
    return re.sub(r"[^A-Za-z0-9_-]+", "_", s).strip("_")[:40] or "fichero"


def unit_dir(work: Path, u: dict) -> Path:
    return work / "unidades" / "{}_{}".format(u["id"], slug(Path(u["name"]).stem))


def read_text_file(path: Path) -> tuple:
    raw = path.read_bytes()
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return raw.decode(enc), enc
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace"), "utf-8 (con sustituciones)"


# ---------------------------------------------------------------- markitdown y su limpieza
def markitdown_text(path: Path) -> str:
    try:
        from markitdown import MarkItDown
    except ImportError:
        sys.exit("error: falta markitdown. Ejecuta desde la raíz del repo: python3 scripts/setup_env.py")
    return MarkItDown(enable_plugins=False).convert(str(path)).text_content or ""


def comment_safe(s: str) -> str:
    return re.sub(r"-{2,}", "-", s).replace(">", "")


def clean_markitdown(md: str, ext: str) -> tuple:
    """Quita lo que markitdown añade y no está en el original. Devuelve (md, nº de imágenes referenciadas)."""
    if ext == ".pptx":
        md = re.sub(r"(?m)^### Notes:\s*$", "<!-- notas del orador -->", md)
        md = re.sub(r"<!-- Slide number: (\d+) -->", r"<!-- diapositiva \1 -->", md)
    count = 0

    def img(m):
        nonlocal count
        count += 1
        alt, src = m.group(1).strip(), m.group(2).strip()
        src = "" if src.startswith("data:") else src
        return "<!-- imagen{}{} -->".format(": " + comment_safe(alt) if alt else "", " ({})".format(comment_safe(src)) if src else "")
    md = re.sub(r"!\[([^\]]*)\]\(([^)]*)\)", img, md)
    md = drop_empty_table_headers(md)
    return md.strip() + "\n", count


def drop_empty_table_headers(md: str) -> str:
    """markitdown (docx) pone una fila de cabecera vacía «|  |  |» y la primera fila real debajo: se sube esa fila."""
    lines = md.split("\n")
    out, i = [], 0
    while i < len(lines):
        if (i + 2 < len(lines) and re.fullmatch(r"\|(\s*\|)+\s*", lines[i]) and re.fullmatch(r"\|(\s*:?-+:?\s*\|)+\s*", lines[i + 1])
                and lines[i + 2].startswith("|")):
            out.extend([lines[i + 2], lines[i + 1]])
            i += 3
            continue
        out.append(lines[i])
        i += 1
    return "\n".join(out)


def fenced(text: str, lang: str) -> str:
    longest = max((len(m) for m in re.findall(r"`+", text)), default=0)
    fence = "`" * max(3, longest + 1)
    return "{}{}\n{}\n{}\n".format(fence, lang, text.rstrip("\n"), fence)


# ---------------------------------------------------------------- hojas de cálculo
def cell_str(v, fmt: str = "General") -> str:
    if v is None:
        return ""
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, dt.datetime):
        return v.date().isoformat() if v.time() == dt.time(0) else v.isoformat(sep=" ")
    if isinstance(v, (dt.date, dt.time)):
        return v.isoformat()
    if isinstance(v, (int, float)):
        fmt = fmt or "General"
        m = re.search(r"0(?:\.(0+))?\s*%", fmt)
        if m:  # porcentaje: Excel muestra el valor ×100
            dec = len(m.group(1) or "")
            return "{:.{}f}%".format(v * 100, dec)
        m = re.search(r"0\.(0+)(?![0#]*%)", fmt)
        if m and isinstance(v, float):  # decimales fijos (0.00, #,##0.00)
            return "{:.{}f}".format(v, len(m.group(1)))
        if isinstance(v, float) and v.is_integer():
            return str(int(v))
        return format(v, ".15g") if isinstance(v, float) else str(v)
    return str(v)


def md_table(rows: List[List[str]]) -> str:
    rows = [[re.sub(r"\s+", " ", c).strip() for c in r] for r in rows]
    while rows and not any(rows[-1]):
        rows.pop()
    if not rows:
        return ""
    ncols = max(len(r) for r in rows)
    while ncols and not any(len(r) >= ncols and r[ncols - 1] for r in rows):
        ncols -= 1
    if not ncols:
        return ""
    rows = [(r + [""] * ncols)[:ncols] for r in rows]
    esc = [[escape_line(c) if c else "" for c in r] for r in rows]
    out = ["| " + " | ".join(esc[0]) + " |", "|" + "|".join([" --- "] * ncols) + "|"]
    out += ["| " + " | ".join(r) + " |" for r in esc[1:]]
    return "\n".join(out)


def sheet_xlsx(path: Path, warnings: List[str]) -> str:
    try:
        import openpyxl
    except ImportError:
        sys.exit("error: falta openpyxl. Ejecuta desde la raíz del repo: python3 scripts/setup_env.py")
    wb = openpyxl.load_workbook(str(path), data_only=True)
    wf = openpyxl.load_workbook(str(path), data_only=False)
    parts = []
    for ws in wb.worksheets:
        rows, no_value = [], []
        for row in ws.iter_rows():
            vals = []
            for c in row:
                v = c.value
                if v is None:
                    f = wf[ws.title][c.coordinate].value
                    if isinstance(f, str) and f.startswith("="):
                        no_value.append(c.coordinate)
                vals.append(cell_str(v, c.number_format))
            rows.append(vals)
        if no_value:
            warnings.append("hoja «{}»: fórmulas sin valor calculado guardado (quedan vacías): {}".format(ws.title, ", ".join(no_value[:20])))
        table = md_table(rows)
        head = "## " + escape_line(ws.title)
        hidden = " <!-- hoja oculta -->" if ws.sheet_state != "visible" else ""
        parts.append(head + hidden + "\n\n" + (table or "<!-- hoja vacía -->"))
    return "\n\n".join(parts) + "\n"


def sheet_xls(path: Path, warnings: List[str]) -> str:
    try:
        import xlrd
    except ImportError:
        sys.exit("error: falta xlrd. Ejecuta desde la raíz del repo: python3 scripts/setup_env.py")
    book = xlrd.open_workbook(str(path))
    parts = []
    for sh in book.sheets():
        rows = []
        for r in range(sh.nrows):
            vals = []
            for c in range(sh.ncols):
                cell = sh.cell(r, c)
                if cell.ctype == xlrd.XL_CELL_DATE:
                    vals.append(cell_str(xlrd.xldate_as_datetime(cell.value, book.datemode)))
                elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                    vals.append(cell_str(bool(cell.value)))
                elif cell.ctype == xlrd.XL_CELL_ERROR:
                    vals.append(xlrd.error_text_from_code.get(cell.value, "#ERROR"))
                else:
                    vals.append(cell_str(cell.value))
            rows.append(vals)
        table = md_table(rows)
        parts.append("## " + escape_line(sh.name) + "\n\n" + (table or "<!-- hoja vacía -->"))
    warnings.append("formato .xls: se escriben los valores guardados, sin formato de número (decimales, %, moneda)")
    return "\n\n".join(parts) + "\n"


def sheet_csv(path: Path, warnings: List[str]) -> str:
    text, enc = read_text_file(path)
    if enc != "utf-8-sig":
        warnings.append("codificación detectada: {}".format(enc))
    if path.suffix.lower() == ".tsv":
        delim = "\t"
    else:
        try:
            delim = csv.Sniffer().sniff(text[:20000], delimiters=",;\t|").delimiter
        except csv.Error:
            delim = ","
    rows = list(csv.reader(io.StringIO(text), delimiter=delim))
    return (md_table(rows) or "<!-- CSV vacío -->") + "\n"


# ---------------------------------------------------------------- PDF
def render_pages(pdf: Path, out: Path, dpi: int, pages: Optional[List[int]] = None, prefix: str = "pag") -> List[str]:
    try:
        import pypdfium2 as pdfium
    except ImportError:
        sys.exit("error: falta pypdfium2. Ejecuta desde la raíz del repo: python3 scripts/setup_env.py")
    out.mkdir(parents=True, exist_ok=True)
    doc = pdfium.PdfDocument(str(pdf))
    width = len(str(len(doc)))
    done = []
    for i in range(len(doc)):
        if pages and i + 1 not in pages:
            continue
        name = "{}_{:0{}d}.png".format(prefix, i + 1, max(width, 3))
        target = out / name
        if not target.exists():
            doc[i].render(scale=dpi / 72).to_pil().save(target)
        done.append(name)
    doc.close()
    return done


# ---------------------------------------------------------------- segundo lector
def second_reader(path: Path) -> Optional[str]:
    ext = path.suffix.lower()
    try:
        if ext == ".pdf":
            from pypdf import PdfReader
            return "\n".join((p.extract_text() or "") for p in PdfReader(str(path)).pages)
        if ext == ".docx":
            import docx
            d = docx.Document(str(path))
            out: List[str] = []

            def walk(container):
                for block in container.iter_inner_content():
                    if hasattr(block, "rows"):
                        for row in block.rows:
                            for cell in row.cells:
                                walk(cell)
                    else:
                        out.append(block.text)
            walk(d)
            return "\n".join(out)
        if ext == ".pptx":
            from pptx import Presentation
            prs = Presentation(str(path))
            out = []

            def shapes(coll):
                for sh in coll:
                    if sh.shape_type == 6 and hasattr(sh, "shapes"):  # grupo
                        shapes(sh.shapes)
                    if getattr(sh, "has_text_frame", False) and sh.has_text_frame:
                        out.append(sh.text_frame.text)
                    if getattr(sh, "has_table", False) and sh.has_table:
                        out.extend(c.text for r in sh.table.rows for c in r.cells)
            for s in prs.slides:
                shapes(s.shapes)
                if s.has_notes_slide:
                    out.append(s.notes_slide.notes_text_frame.text)
            return "\n".join(out)
        if ext in (".html", ".htm", ".xhtml"):
            from bs4 import BeautifulSoup
            text, _ = read_text_file(path)
            soup = BeautifulSoup(text, "html.parser")
            for t in soup(["script", "style", "noscript", "head"]):
                t.decompose()
            return soup.get_text("\n")
    except Exception as e:  # el segundo lector es solo un control: si falla, se avisa y se sigue
        return "__ERROR__ {}".format(e)
    return None


def docx_headers_footers(path: Path) -> str:
    try:
        import docx
        d = docx.Document(str(path))
        seen, out = set(), []
        for s in d.sections:
            for part in (s.header, s.footer, s.first_page_header, s.first_page_footer):
                t = "\n".join(p.text for p in part.paragraphs).strip()
                if t and t not in seen:
                    seen.add(t)
                    out.append(t)
        return "\n".join(out)
    except Exception:
        return ""


def compare_second(bruto_md: str, second: str, list_numbers: bool) -> dict:
    # En PDF los números de las listas son texto; en docx/pptx/html son numeración automática y el segundo lector no los ve
    a, b = Counter(content_words(bruto_md, list_numbers)), Counter(text_words(second))
    missing = b - a
    extra = a - b
    return {"words_bruto": sum(a.values()), "words_second": sum(b.values()),
            "missing": [w for w, _ in missing.most_common(40)], "missing_count": sum(missing.values()),
            "extra": [w for w, _ in extra.most_common(40)], "extra_count": sum(extra.values())}


def embedded_images(path: Path, out: Path) -> List[str]:
    if path.suffix.lower() not in (".docx", ".pptx", ".xlsx", ".xlsm", ".epub"):
        return []
    names = []
    with zipfile.ZipFile(path) as z:
        for info in z.infolist():
            p = Path(info.filename)
            if p.suffix.lower() in IMAGE_EXTS and ("media" in p.parts or path.suffix.lower() == ".epub"):
                out.mkdir(parents=True, exist_ok=True)
                target = out / "{:02d}_{}".format(len(names) + 1, p.name)
                target.write_bytes(z.read(info))
                names.append(target.name)
    return names


# ---------------------------------------------------------------- unidades
def office_convert(u: dict, udir: Path) -> Path:
    exe = shutil.which("soffice") or shutil.which("libreoffice")
    if not exe:
        sys.exit("error: {} necesita LibreOffice (soffice) y no está instalado".format(u["name"]))
    outdir = udir / "convertido"
    outdir.mkdir(parents=True, exist_ok=True)
    target = outdir / (Path(u["name"]).stem + "." + u["convert_to"])
    if not target.exists():
        filt = {"docx": "docx:MS Word 2007 XML", "pptx": "pptx", "xlsx": "xlsx"}[u["convert_to"]]
        r = subprocess.run([exe, "--headless", "--convert-to", filt, "--outdir", str(outdir), u["path"]],
                           capture_output=True, text=True, timeout=600)
        if not target.exists():
            sys.exit("error: LibreOffice no pudo convertir {}: {}".format(u["name"], (r.stderr or r.stdout).strip()[:300]))
    return target


def process_unit(u: dict, work: Path, force: bool) -> dict:
    udir = unit_dir(work, u)
    state_path = udir / "estado.json"
    if state_path.is_file() and not force:
        st = json.loads(state_path.read_text())
        st["reused"] = True
        return st
    if udir.exists() and force:
        for sub in ("bruto.md", "segundo_lector.txt", "estado.json"):
            (udir / sub).unlink(missing_ok=True)
    udir.mkdir(parents=True, exist_ok=True)
    src = Path(u["path"])
    st = {"id": u["id"], "name": u["name"], "kind": u["kind"], "dir": str(udir), "warnings": [], "needs": []}
    if u["kind"] == "office":
        src = office_convert(u, udir)
        st["converted"] = str(src)
        st["warnings"].append("convertido con LibreOffice a {}: revisa que no se haya perdido nada".format(src.name))
    ext = src.suffix.lower()
    md: Optional[str] = None

    if u["kind"] in ("audio",):
        st["needs"].append("audio: ejecutar process_audio.py")
    elif u["kind"] == "image":
        st["needs"].append("imagen: transcribir con img-to-text en trabajo/imagenes")
    elif u.get("scanned"):
        pages = render_pages(src, udir / "paginas", SCAN_DPI)
        st["pages_dir"] = str(udir / "paginas")
        st["needs"].append("PDF escaneado: transcribir las {} página(s) de paginas/ con img-to-text en {}/img".format(len(pages), udir))
    elif u["kind"] == "text":
        text, enc = read_text_file(src)
        if enc != "utf-8-sig":
            st["warnings"].append("codificación detectada: {}".format(enc))
        if ext in (".md", ".markdown"):
            md = text
        else:
            md = escape_text(text)
            if text_words(text) != content_words(md):  # el escape no puede cambiar palabras
                sys.exit("error interno: el escape de {} cambia palabras; documenta el caso en references/errores_conocidos.md".format(src))
    elif ext in (".xlsx", ".xlsm", ".xls", ".csv", ".tsv"):
        md = {".xls": sheet_xls, ".csv": sheet_csv, ".tsv": sheet_csv}.get(ext, sheet_xlsx)(src, st["warnings"])
    else:
        raw = markitdown_text(src)
        if ext in (".json", ".jsonl", ".xml", ".rss", ".atom"):
            md = fenced(raw if raw.strip() else read_text_file(src)[0], ext.lstrip(".").replace("jsonl", "json"))
        else:
            md, n_img = clean_markitdown(raw, ext)
            if n_img:
                st["warnings"].append("{} imagen(es) referenciadas: quedan como comentario <!-- imagen … --> (no se lee su texto)".format(n_img))
        if ext == ".pdf":
            if u.get("textless_pages"):
                render_pages(src, udir / "paginas", SCAN_DPI, u["textless_pages"])
                st["pages_dir"] = str(udir / "paginas")
                st["needs"].append("PDF con páginas SIN TEXTO {}: transcribirlas con img-to-text e insertarlas en su sitio".format(u["textless_pages"]))
            st["preview_dir"] = str(udir / "vista")
            render_pages(src, udir / "vista", PREVIEW_DPI)
        if ext == ".docx":
            hf = docx_headers_footers(src)
            if hf:
                st["warnings"].append("el documento tiene cabeceras/pies de página que markitdown no incluye: «{}»".format(hf[:150].replace("\n", " / ")))

    if md is not None:
        if not md.strip():
            st["warnings"].append("la extracción está VACÍA: ¿es un documento escaneado o protegido?")
        (udir / "bruto.md").write_text(md, encoding="utf-8")
        st["bruto"] = str(udir / "bruto.md")
        second = second_reader(src)
        if second is not None:
            if second.startswith("__ERROR__"):
                st["warnings"].append("el segundo lector falló: " + second[10:])
            else:
                (udir / "segundo_lector.txt").write_text(second, encoding="utf-8")
                cmp = compare_second(md, second, list_numbers=ext == ".pdf")
                st["second_reader"] = cmp
                if cmp["missing_count"]:
                    st["warnings"].append("OMITIDO? {} palabra(s) del segundo lector no están en bruto.md: {}".format(
                        cmp["missing_count"], ", ".join(cmp["missing"][:20])))
                if cmp["extra_count"]:
                    st["warnings"].append("AÑADIDO? {} palabra(s) de bruto.md no las ve el segundo lector: {}".format(
                        cmp["extra_count"], ", ".join(cmp["extra"][:20])))
    imgs = embedded_images(src, udir / "incrustadas")
    if imgs:
        st["embedded_images"] = imgs
        st["needs"].append("{} imagen(es) incrustada(s) en incrustadas/: PREGUNTAR al usuario si se transcribe su texto".format(len(imgs)))
    state_path.write_text(json.dumps(st, ensure_ascii=False, indent=2) + "\n")
    return st


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("work_dir", type=Path, help="carpeta trabajo/ creada por inventory.py")
    ap.add_argument("--unit", action="append", help="solo esta unidad (u01…); repetible")
    ap.add_argument("--force", action="store_true", help="rehacer la extracción (no toca final.md)")
    args = ap.parse_args()

    inv_path = args.work_dir / "inventario.json"
    if not inv_path.is_file():
        sys.exit("error: falta {} (ejecuta inventory.py)".format(inv_path))
    inv = json.loads(inv_path.read_text())
    if not inv.get("formats"):
        sys.exit("error: el inventario no tiene formatos de salida: pregunta al usuario y reejecuta inventory.py con --formats")
    units = [u for u in inv["units"] if not args.unit or u["id"] in args.unit]
    summary = []
    for u in units:
        st = process_unit(u, args.work_dir, args.force)
        summary.append(st)
        print("\n=== {} · {} ({}){} ===".format(u["id"], u["name"], u["kind"], " [reutilizado]" if st.get("reused") else ""))
        if st.get("bruto"):
            words = len(content_words(Path(st["bruto"]).read_text(encoding="utf-8")))
            print("  bruto.md: {} palabras → {}".format(words, st["bruto"]))
        if st.get("preview_dir"):
            print("  vistas de página (para comparar el formato): {}".format(st["preview_dir"]))
        for w in st["warnings"]:
            print("  AVISO:", w)
        for n in st["needs"]:
            print("  PENDIENTE:", n)
    (args.work_dir / "conversion.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2) + "\n")
    print("\nSiguiente: revisar el formato de cada bruto.md → final.md (references/reglas_formato.md) y completar lo PENDIENTE.")


if __name__ == "__main__":
    main()
