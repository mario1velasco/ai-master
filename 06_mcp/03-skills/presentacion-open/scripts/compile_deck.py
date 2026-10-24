#!/usr/bin/env python3
"""Compile editable PPTX decks using the public python-pptx API."""
from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path
from typing import Any, Dict, List, Mapping, Optional

from pptx import Presentation
from pptx.util import Inches

from pptx_components import (
    COLUMN_GAP_IN, DEFAULT_MARGIN_IN, DEFAULT_SLIDE_HEIGHT_IN, DEFAULT_SLIDE_WIDTH_IN,
    DEFAULT_THEME, EMU_PER_INCH, ITEM_GAP_IN, MAX_CONTAINER_CHARS,
    TEXT_FRAME_MARGIN_IN, add_chart, add_element, add_image, add_table, add_textbox, apply_text_frame,
    enforce_container, estimated_text_height, fail, finite_number, line_count,
    parse_color, partition_by_height, plain_text, split_long_text,
)

def _set_shape_text(shape: Any, content: Any, theme: Mapping[str, Any], preserve_format: bool = True) -> None:
    if not getattr(shape, "has_text_frame", False):
        fail("Selected placeholder or shape has no text frame")
    if isinstance(content, str) and preserve_format:
        enforce_container(content, "placeholder text")
        shape.text_frame.text = content
    else:
        apply_text_frame(shape.text_frame, content, theme, content.get("style", {}) if isinstance(content, dict) else {}, use_theme=not preserve_format)


def _placeholder(slide: Any, index: Any) -> Any:
    try:
        wanted = int(index)
    except (TypeError, ValueError):
        fail("placeholder keys must be numeric placeholder indices")
    for shape in slide.placeholders:
        if shape.placeholder_format.idx == wanted:
            return shape
    fail("Slide has no placeholder with idx {}".format(wanted))


def _prepare_auto_blocks(values: List[Any], location: str, bullet: bool) -> List[Any]:
    blocks: List[Any] = []
    for index, value in enumerate(values):
        text = enforce_container(value, "{}[{}]".format(location, index), maximum=10000)
        if isinstance(value, str):
            blocks.extend(split_long_text(text, "{}[{}]".format(location, index)))
        else:
            if isinstance(value, dict) and isinstance(value.get("paragraphs"), list):
                paragraph_count = max(1, len(value["paragraphs"]))
            elif isinstance(value, list):
                paragraph_count = max(1, len(value))
            else:
                paragraph_count = max(1, len(text.splitlines()))
            rendered_length = len(text) + (2 * paragraph_count if bullet else 0)
            if rendered_length > MAX_CONTAINER_CHARS:
                fail("{}[{}] is rich text longer than 180 characters; split it into shorter blocks".format(location, index))
            blocks.append(value)
    return blocks


def _layout_index(presentation: Any, slide_data: Mapping[str, Any], has_template: bool) -> int:
    if "layout_index" in slide_data:
        index = int(slide_data["layout_index"])
        if index < 0 or index >= len(presentation.slide_layouts):
            fail("layout_index is outside the presentation's available layouts")
        return index
    requested = slide_data.get("layout")
    legacy_layout_modes = {"auto", "single", "two_columns", "three_columns"}
    if isinstance(requested, str) and requested in legacy_layout_modes:
        requested = None
    if requested is not None:
        if isinstance(requested, int):
            if requested < 0 or requested >= len(presentation.slide_layouts):
                fail("layout index is outside the presentation's available layouts")
            return requested
        for index, layout in enumerate(presentation.slide_layouts):
            if layout.name == str(requested):
                return index
        fail("No slide layout named '{}'".format(requested))
    if has_template:
        for index, layout in enumerate(presentation.slide_layouts):
            if layout.name.lower() in ("title and content", "title only", "title slide"):
                return index
    return min(6, len(presentation.slide_layouts) - 1)


def set_slide_background(slide: Any, color: str) -> None:
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = parse_color(color, "theme.background_color")


def set_title(slide: Any, title: Any, theme: Mapping[str, Any], y: float, width: float, margin: float, preserve_template: bool) -> float:
    text = enforce_container(title, "slide title")
    title_height = max(0.58, estimated_text_height(text, width, float(theme.get("title_pt", 32))))
    title_shape = getattr(slide.shapes, "title", None)
    if preserve_template and title_shape is not None:
        title_shape.text_frame.text = text
    else:
        add_textbox(slide, text, margin + 0.1, y, width, title_height, theme, {"font_size_pt": theme.get("title_pt", 32), "color": theme.get("title_color", "17324D"), "bold": True}, "Slide title")
    return y + title_height + 0.14


def add_standard_body(slide: Any, slide_data: Mapping[str, Any], deck_dir: Path, slide_width: float, slide_height: float, theme: Mapping[str, Any], body_top: float) -> None:
    margin = float(theme["safe_margin_inches"])
    sources = slide_data.get("sources", [])
    source_reserve = 0.0
    body_bottom = slide_height - margin - source_reserve
    body_height = body_bottom - body_top
    if body_height <= 0.6:
        fail("Title leaves too little room for slide content")
    body_width = slide_width - 2 * (margin + 0.1)
    body_color = theme.get("body_color", "243B53")
    font_size = float(theme.get("body_pt", 20))
    font_name = theme.get("font_name", "Roboto")

    if "elements" in slide_data:
        if any(key in slide_data for key in ("items", "columns", "table", "chart", "image")):
            fail("Use elements for custom content, or use the automatic items/table/chart/image fields, not both")
        elements = slide_data.get("elements")
        if not isinstance(elements, list):
            fail("slide.elements must be a list")
        for index, element in enumerate(elements):
            add_element(slide, element, deck_dir, theme, index)
        return

    table_spec = slide_data.get("table")
    chart_spec = slide_data.get("chart")
    if table_spec is not None or chart_spec is not None:
        if sum(value is not None for value in (table_spec, chart_spec)) > 1 or any(key in slide_data for key in ("items", "columns", "image")):
            fail("Automatic table/chart layout is exclusive; use elements to combine content")
        object_spec = dict(table_spec if table_spec is not None else chart_spec)
        object_spec.setdefault("x", margin + 0.1)
        object_spec.setdefault("y", body_top)
        object_spec.setdefault("width", body_width)
        object_spec.setdefault("height", body_height)
        if table_spec is not None:
            add_table(slide, object_spec, theme, "Editable table")
        else:
            add_chart(slide, object_spec, theme, "Editable chart")
        return

    image_spec = slide_data.get("image")
    items = slide_data.get("items", [])
    raw_columns = slide_data.get("columns")
    bullet = bool(slide_data.get("bullet", True))
    if isinstance(items, str):
        items = [items]
    if raw_columns is not None:
        if not isinstance(raw_columns, list) or not 1 <= len(raw_columns) <= 3:
            fail("columns must contain one to three lists")
        if image_spec is not None and len(raw_columns) > 1:
            fail("Automatic image layout supports one text column; use elements for a custom layout")
        columns: List[List[Any]] = []
        for col_index, raw_column in enumerate(raw_columns):
            if not isinstance(raw_column, list):
                fail("columns[{}] must be a list".format(col_index))
            columns.append(_prepare_auto_blocks(raw_column, "columns[{}]".format(col_index), bullet))
    else:
        if not isinstance(items, list):
            fail("items must be a list of text strings or rich-text objects")
        flat = _prepare_auto_blocks(items, "items", bullet)
        if image_spec is not None:
            image_width = min(4.8, body_width * 0.38)
            content_width = body_width - image_width - 0.35
        else:
            image_width = 0.0
            content_width = body_width
        legacy_layout = slide_data.get("layout")
        default_layout_mode = legacy_layout if isinstance(legacy_layout, str) and legacy_layout in ("auto", "single", "two_columns", "three_columns") else "auto"
        layout_mode = str(slide_data.get("layout_mode", default_layout_mode))
        choices = {"single": [1], "two_columns": [2], "three_columns": [3], "auto": [1, 2, 3]}
        if layout_mode not in choices:
            fail("layout_mode must be auto, single, two_columns, or three_columns")
        columns = []
        for count in choices[layout_mode]:
            if count > max(1, len(flat)):
                continue
            gap = COLUMN_GAP_IN * (count - 1)
            col_width = (content_width - gap) / count
            # Greedy contiguous split by estimated rendered height, preserving reading order.
            groups = partition_by_height(flat, count, col_width, font_size, bullet)
            heights = [sum(estimated_text_height(item, col_width, font_size, bullet) + ITEM_GAP_IN for item in group) - (ITEM_GAP_IN if group else 0) for group in groups]
            if max(heights, default=0) <= body_height:
                columns = groups
                break
        if flat and not columns:
            fail("Body does not fit at the configured font size; shorten it or split it across slides")

    col_count = max(1, len(columns))
    content_width = body_width - (min(4.8, body_width * 0.38) + 0.35 if image_spec is not None else 0)
    gap = COLUMN_GAP_IN * (col_count - 1)
    col_width = (content_width - gap) / col_count
    for col_index, column in enumerate(columns):
        x = margin + 0.1 + col_index * (col_width + COLUMN_GAP_IN)
        y = body_top
        for item_index, text in enumerate(column):
            height = max(0.32, estimated_text_height(text, col_width, font_size, bullet))
            add_textbox(slide, text, x, y, col_width, height, theme, {"font_name": font_name, "font_size_pt": font_size, "color": body_color}, "Body {}.{}".format(col_index + 1, item_index + 1), bullet=bullet)
            y += height + ITEM_GAP_IN
    if image_spec is not None:
        image_width = min(4.8, body_width * 0.38)
        add_image(slide, {"path": image_spec.get("path") if isinstance(image_spec, dict) else image_spec, "x": slide_width - margin - image_width, "y": body_top, "width": image_width, "height": body_height, "fit": image_spec.get("fit", "contain") if isinstance(image_spec, dict) else "contain"}, deck_dir, "Slide image")


def add_slide_notes(slide: Any, notes: Any, sources: Any) -> None:
    notes_list: List[str] = []
    if notes:
        notes_list.append(str(notes))
    if sources:
        if not isinstance(sources, list):
            fail("sources must be a list of strings")
        source_lines = [enforce_container(source, "source") for source in sources]
        notes_list.append("Fuentes:\n" + "\n".join(source_lines))
    if not notes_list:
        return
    notes_slide = slide.notes_slide
    text_frame = notes_slide.notes_text_frame
    if text_frame is None:
        fail("This slide layout has no notes text placeholder; cannot store notes/sources")
    text_frame.text = "\n\n".join(notes_list)


def fill_slide_placeholders(slide: Any, slide_data: Mapping[str, Any], theme: Mapping[str, Any]) -> None:
    placeholders = slide_data.get("placeholders", {})
    if not isinstance(placeholders, dict):
        fail("slide.placeholders must be an object keyed by placeholder index")
    for index, content in placeholders.items():
        shape = _placeholder(slide, index)
        _set_shape_text(shape, content, theme, preserve_format=isinstance(content, str))
    named = slide_data.get("replace_shapes", {})
    if not isinstance(named, dict):
        fail("slide.replace_shapes must map exact shape names to text")
    if named:
        by_name = {shape.name: shape for shape in slide.shapes}
        for shape_name, content in named.items():
            if shape_name not in by_name:
                fail("No existing shape named '{}'".format(shape_name))
            _set_shape_text(by_name[shape_name], content, theme, preserve_format=isinstance(content, str))


def compile_one_slide(presentation: Any, slide: Any, slide_data: Mapping[str, Any], deck_dir: Path, theme: Mapping[str, Any], is_new: bool, use_template_style: bool) -> None:
    width = float(presentation.slide_width) / EMU_PER_INCH
    height = float(presentation.slide_height) / EMU_PER_INCH
    margin = float(theme["safe_margin_inches"])
    if is_new:
        set_slide_background(slide, str(theme.get("background_color", "FFFFFF")))
    fill_slide_placeholders(slide, slide_data, theme)
    title = slide_data.get("title")
    kind = str(slide_data.get("kind", "content")).lower()
    if kind not in ("cover", "content", "existing"):
        fail("slide.kind must be cover, content, or existing")

    if kind == "cover":
        title_pt = float(theme.get("cover_title_pt", 42))
        cover_image = slide_data.get("image")
        body_width = width - 2 * (margin + 0.1)
        if cover_image is not None:
            body_width = min(body_width, width * 0.55)
        if title:
            title_text = enforce_container(title, "cover title")
            title_height = max(1.0, estimated_text_height(title_text, body_width, title_pt) + 0.12)
            title_shape = getattr(slide.shapes, "title", None) if use_template_style else None
            if title_shape is not None:
                title_shape.text_frame.text = title_text
            else:
                add_textbox(slide, title_text, margin + 0.1, max(1.2, height * 0.28), body_width, title_height, theme, {"font_size_pt": title_pt, "color": theme.get("title_color", "17324D"), "bold": True}, "Cover title")
        if slide_data.get("subtitle"):
            subtitle = enforce_container(slide_data["subtitle"], "cover subtitle")
            add_textbox(slide, subtitle, margin + 0.1, max(1.2, height * 0.28) + 1.1, body_width, 0.7, theme, {"font_size_pt": float(theme.get("body_pt", 20)), "color": theme.get("body_color", "243B53")}, "Cover subtitle")
        if cover_image is not None:
            image_width = min(4.8, width * 0.38)
            add_image(slide, {"path": cover_image.get("path") if isinstance(cover_image, dict) else cover_image, "x": width - margin - image_width, "y": margin + 0.1, "width": image_width, "height": height - 2 * margin - 0.2, "fit": cover_image.get("fit", "contain") if isinstance(cover_image, dict) else "contain"}, deck_dir, "Cover image")
        if "elements" in slide_data:
            for index, element in enumerate(slide_data["elements"]):
                add_element(slide, element, deck_dir, theme, index)
    elif kind == "existing":
        if not slide_data.get("target_slide"):
            fail("kind 'existing' requires target_slide")
    else:
        body_top = margin + 0.05
        if title:
            body_top = set_title(slide, title, theme, body_top, width - 2 * (margin + 0.1), margin, preserve_template=use_template_style)
        if slide_data.get("subtitle"):
            subtitle = enforce_container(slide_data["subtitle"], "slide subtitle")
            add_textbox(slide, subtitle, margin + 0.1, body_top, width - 2 * (margin + 0.1), 0.52, theme, {"font_size_pt": float(theme.get("subtitle_pt", 20)), "color": theme.get("body_color", "243B53")}, "Slide subtitle")
            body_top += 0.62
        if any(key in slide_data for key in ("elements", "items", "columns", "table", "chart", "image")):
            add_standard_body(slide, slide_data, deck_dir, width, height, theme, body_top)
    add_slide_notes(slide, slide_data.get("notes"), slide_data.get("sources"))


def _template_path(deck: Mapping[str, Any], deck_path: Path) -> Optional[Path]:
    spec = deck.get("template")
    if spec is None:
        return None
    raw = spec.get("path") if isinstance(spec, dict) else spec
    if not isinstance(raw, str) or not raw.strip():
        fail("template must be a path or object containing path")
    path = Path(raw)
    if not path.is_absolute():
        path = (deck_path.parent / path).resolve()
    if not path.is_file():
        fail("Template PPTX not found: {}".format(path))
    return path


def _font_policy(deck: Mapping[str, Any], theme: Mapping[str, Any]) -> None:
    policy = deck.get("font_policy")
    if policy is None:
        return
    if not isinstance(policy, dict):
        fail("font_policy must be an object")
    if policy.get("basis") not in ("design", "user_request", "reference"):
        fail("font_policy.basis must be design, user_request, or reference")
    families = policy.get("families")
    if not isinstance(families, list) or not families or any(not isinstance(x, str) or not x for x in families):
        fail("font_policy.families must be a non-empty list of font family names")
    if theme.get("font_name") not in families:
        fail("theme.font_name must be included in font_policy.families")


def compile_deck(deck_path: Path, output_path: Optional[Path]) -> Path:
    with deck_path.open("r", encoding="utf-8") as handle:
        deck = json.load(handle)
    if not isinstance(deck, dict):
        fail("deck.json root must be an object")
    version = deck.get("version", 1)
    if version not in (1, 2):
        fail("Only deck schema versions 1 and 2 are supported")
    slides = deck.get("slides")
    if not isinstance(slides, list) or not slides:
        fail("deck.json must contain a non-empty slides list")
    for index, slide_data in enumerate(slides):
        if not isinstance(slide_data, dict):
            fail("slides[{}] must be an object".format(index))

    raw_theme = deck.get("theme", {})
    if not isinstance(raw_theme, dict):
        fail("theme must be an object")
    theme = dict(DEFAULT_THEME)
    theme.update(raw_theme)
    margin = finite_number(theme.get("safe_margin_inches"), "theme.safe_margin_inches", 0.0)
    if margin < 0.3:
        fail("theme.safe_margin_inches must be at least 0.3")
    for key in ("title_color", "body_color", "accent_color", "background_color", "muted_color"):
        parse_color(theme[key], "theme." + key)
    for key in ("title_pt", "body_pt", "cover_title_pt"):
        finite_number(theme[key], "theme." + key, 1)
    _font_policy(deck, theme)

    template_path = _template_path(deck, deck_path)
    if template_path:
        presentation = Presentation(str(template_path))
        template_slide_count = len(presentation.slides)
        template_size = (float(presentation.slide_width) / EMU_PER_INCH, float(presentation.slide_height) / EMU_PER_INCH)
        slide_size = deck.get("slide_size")
        if slide_size is not None:
            if not isinstance(slide_size, dict):
                fail("slide_size must be an object")
            requested = (finite_number(slide_size.get("width_in"), "slide_size.width_in", 1), finite_number(slide_size.get("height_in"), "slide_size.height_in", 1))
            if any(abs(a - b) > 0.01 for a, b in zip(requested, template_size)):
                fail("Do not resize a supplied template; remove slide_size or match its dimensions")
        slide_width, slide_height = template_size
    else:
        base_path = Path(__file__).resolve().parent.parent / "assets" / "base_theme.pptx"
        base_seed = False
        if base_path.is_file():
            presentation = Presentation(str(base_path))
            base_seed = (len(presentation.slides) == 1 and len(presentation.slides[0].shapes) == 0)
            if not base_seed:
                fail("assets/base_theme.pptx must contain one blank slide when used as the default seed")
        else:
            presentation = Presentation()
        slide_size = deck.get("slide_size", {})
        if not isinstance(slide_size, dict):
            fail("slide_size must be an object")
        slide_width = finite_number(slide_size.get("width_in", DEFAULT_SLIDE_WIDTH_IN), "slide_size.width_in", 1)
        slide_height = finite_number(slide_size.get("height_in", DEFAULT_SLIDE_HEIGHT_IN), "slide_size.height_in", 1)
        if slide_width <= 2 * margin or slide_height <= 2 * margin:
            fail("Slide dimensions are too small for the configured safe margin")
        presentation.slide_width = Inches(slide_width)
        presentation.slide_height = Inches(slide_height)
        template_slide_count = 0

    seed_slide_available = bool(not template_path and base_path.is_file() and len(presentation.slides) == 1) if not template_path else False
    for index, slide_data in enumerate(slides):
        if slide_data.get("target_slide") is not None:
            target = int(slide_data["target_slide"]) - 1
            if target < 0 or target >= len(presentation.slides):
                fail("slides[{}].target_slide must refer to an existing one-based slide".format(index))
            slide = presentation.slides[target]
            compile_one_slide(presentation, slide, slide_data, deck_path.parent, theme, False, True)
            continue
        layout_index = _layout_index(presentation, slide_data, bool(template_path))
        layout = presentation.slide_layouts[layout_index]
        if seed_slide_available:
            slide = presentation.slides[0]
            seed_slide_available = False
        else:
            slide = presentation.slides.add_slide(layout)
        compile_one_slide(presentation, slide, slide_data, deck_path.parent, theme, True, bool(template_path))

    metadata = deck.get("metadata", {})
    if not isinstance(metadata, dict):
        fail("metadata must be an object")
    core = presentation.core_properties
    if "title" in metadata:
        core.title = str(metadata["title"])
    if "subject" in metadata:
        core.subject = str(metadata["subject"])
    if "author" in metadata:
        core.author = str(metadata["author"])
    if "keywords" in metadata:
        core.keywords = str(metadata["keywords"])
    if "comments" in metadata:
        core.comments = str(metadata["comments"])

    requirements = deck.get("requirements", {})
    if not isinstance(requirements, dict):
        fail("requirements must be an object")
    final_count = len(presentation.slides)
    if "total_slides" in requirements and final_count != int(requirements["total_slides"]):
        fail("Output has {} slides; requirements.total_slides is {}".format(final_count, requirements["total_slides"]))
    if "required_native_table_slides" in requirements or "required_native_chart_slides" in requirements:
        from check_overflow import inspect_presentation
        qa_issues = inspect_presentation(presentation, margin, requirements)
        if qa_issues:
            fail("Pre-save QA failed:\n- " + "\n- ".join(qa_issues))

    if output_path is None:
        output_path = deck_path.parent / "out" / "presentation.pptx"
    elif not output_path.is_absolute():
        output_path = (Path.cwd() / output_path).resolve()
    output_path.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(str(output_path))
    return output_path


def main() -> int:
    parser = argparse.ArgumentParser(description="Compile an editable PowerPoint from deck.json using python-pptx")
    parser.add_argument("deck_json", type=Path, help="Input JSON deck specification")
    parser.add_argument("--output", type=Path, default=None, help="Output PPTX path; default: out/presentation.pptx beside deck.json")
    args = parser.parse_args()
    deck_path = args.deck_json.resolve()
    if not deck_path.is_file():
        parser.error("deck.json not found: {}".format(deck_path))
    try:
        result = compile_deck(deck_path, args.output)
    except (OSError, ValueError, TypeError, KeyError, NotImplementedError, json.JSONDecodeError) as error:
        print("ERROR: {}".format(error), file=sys.stderr)
        return 1
    print("Created editable presentation: {}".format(result))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
